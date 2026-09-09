"""CSV / Excel renderers.

Both take the *same* section list from :mod:`reports.report_data`, so the two
downloads always carry identical data. Add a section there and it appears in
both files automatically.
"""

import csv
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone

from .report_data import DEFAULT_DAYS, build_report, parse_days  # noqa: F401 (re-exported)

__all__ = ["parse_days", "build_csv", "build_excel", "DEFAULT_DAYS", "render_csv", "render_excel"]


def _filename(prefix, extension):
    stamp = timezone.now().strftime("%Y-%m-%d")
    return f"{prefix}_{stamp}.{extension}"


# ---------------------------------------------------------------------------
# Renderers - these know about file formats, nothing about the report content
# ---------------------------------------------------------------------------

def render_csv(sections, *, filename):
    """Write ``sections`` as one CSV, each section a titled block."""
    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    writer = csv.writer(response)

    for index, section in enumerate(sections):
        if index:
            writer.writerow([])
        writer.writerow([section.title])
        writer.writerow(section.headers)
        for row in section.rows:
            writer.writerow(row)
        if not section.rows:
            writer.writerow(["(no records)"])

    return response


def render_excel(sections, *, filename):
    """Write ``sections`` as one workbook, each section a worksheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    header_fill = PatternFill(start_color="0066CC", end_color="0066CC", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)

    wb = Workbook()
    wb.remove(wb.active)  # start clean; every sheet comes from a section

    for section in sections:
        ws = wb.create_sheet(section.key[:31])
        ws.append(section.headers)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for row in section.rows:
            ws.append(list(row))
        if not section.rows:
            ws.append(["(no records)"])

        ws.freeze_panes = "A2"
        for column in ws.columns:
            longest = 0
            letter = column[0].column_letter
            for cell in column:
                try:
                    longest = max(longest, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[letter].width = min(longest + 2, 50)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


# ---------------------------------------------------------------------------
# The full inventory report
# ---------------------------------------------------------------------------

def build_csv(days=DEFAULT_DAYS):
    return render_csv(
        build_report(days), filename=_filename("CSE_Inventory_Report", "csv")
    )


def build_excel(days=DEFAULT_DAYS):
    return render_excel(
        build_report(days), filename=_filename("CSE_Inventory_Report", "xlsx")
    )


# ---------------------------------------------------------------------------
# The per-room snapshot report
# ---------------------------------------------------------------------------

def _snapshot_sections():
    from .selectors import room_snapshot

    from .report_data import Section

    snapshot = room_snapshot()
    overview = Section(
        "Snapshot Summary",
        "ROOM INVENTORY SNAPSHOT",
        ["Metric", "Value"],
        [
            ["Snapshot taken", snapshot["generated_at"].strftime("%Y-%m-%d %H:%M:%S")],
            ["Rooms", snapshot["room_count"]],
            ["Distinct items", snapshot["item_count"]],
            ["Total units on hand", snapshot["total_quantity"]],
            ["Low stock items", snapshot["low_stock_count"]],
        ],
    )

    detail_rows = []
    for room in snapshot["rooms"]:
        for item in room["items"]:
            detail_rows.append(
                [
                    room["room_name"],
                    room["room_type"] or "",
                    room["location"] or "",
                    item["item_name"],
                    item["category"],
                    item["quantity"],
                    item["unit"],
                    item["min_quantity"],
                    "Low Stock" if item["is_low_stock"] else "OK",
                ]
            )
    detail = Section(
        "By Room",
        "ITEMS BY ROOM",
        ["Room", "Room Type", "Location", "Item", "Category", "Quantity", "Unit", "Min Quantity", "Status"],
        detail_rows,
    )

    totals = Section(
        "Room Totals",
        "ROOM TOTALS",
        ["Room", "Room Type", "Location", "Distinct Items", "Total Units", "Low Stock Items"],
        [
            [
                room["room_name"],
                room["room_type"] or "",
                room["location"] or "",
                room["item_count"],
                room["total_quantity"],
                room["low_stock_count"],
            ]
            for room in snapshot["rooms"]
        ],
    )
    return [overview, totals, detail]


def build_snapshot_csv():
    return render_csv(
        _snapshot_sections(), filename=_filename("CSE_Room_Snapshot", "csv")
    )


def build_snapshot_excel():
    return render_excel(
        _snapshot_sections(), filename=_filename("CSE_Room_Snapshot", "xlsx")
    )
