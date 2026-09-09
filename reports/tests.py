import csv
import io
from datetime import timedelta

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone

from catalog.models import Item, Room
from common.exceptions import DomainError
from reports import exports
from reports.report_data import build_report
from reports.selectors import (
    dashboard_stats,
    low_stock_items,
    reconciliation_report,
    room_snapshot,
    stock_movement_series,
)
from reports.services import reconcile_item
from stock.models import StockTransaction

User = get_user_model()


class ReconciliationReportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("a@x.com", "Admin", role="admin")

    def _item(self, qty, opening):
        return Item.objects.create(
            item_name="Widget", unit="pcs", quantity=qty, opening_quantity=opening
        )

    def test_balanced_item_when_quantity_matches_ledger(self):
        item = self._item(qty=15, opening=10)
        StockTransaction.objects.create(item=item, type="IN", quantity=10)
        StockTransaction.objects.create(item=item, type="OUT", quantity=5)

        row = reconciliation_report()["rows"][0]
        self.assertEqual(row["expected_quantity"], 15)
        self.assertEqual(row["discrepancy"], 0)
        self.assertEqual(row["status"], "balanced")

    def test_mismatch_is_flagged_when_quantity_edited_outside_ledger(self):
        item = self._item(qty=100, opening=10)
        StockTransaction.objects.create(item=item, type="IN", quantity=10)

        report = reconciliation_report()
        row = report["rows"][0]
        self.assertEqual(row["expected_quantity"], 20)
        self.assertEqual(row["discrepancy"], 80)
        self.assertEqual(row["status"], "over")
        self.assertEqual(report["mismatched"], 1)
        self.assertEqual(report["balanced"], 0)

    def test_reconcile_item_posts_adjustment_and_balances(self):
        item = self._item(qty=100, opening=10)
        StockTransaction.objects.create(item=item, type="IN", quantity=10)

        result = reconcile_item(item_id=item.item_id, user=self.user)
        self.assertEqual(result["adjustment"], 80)

        adjust = StockTransaction.objects.get(item=item, type="ADJUST")
        self.assertEqual(adjust.quantity, 80)

        row = reconciliation_report()["rows"][0]
        self.assertEqual(row["discrepancy"], 0)
        self.assertEqual(row["status"], "balanced")

    def test_reconcile_rejects_already_balanced_item(self):
        item = self._item(qty=10, opening=10)
        with self.assertRaises(DomainError):
            reconcile_item(item_id=item.item_id, user=self.user)


def _parse_csv_sections(response):
    """Pull section blocks back out of a rendered CSV: {title: (headers, rows)}."""
    reader = csv.reader(io.StringIO(response.content.decode("utf-8")))
    sections, title, headers, rows = {}, None, None, []
    for raw in reader:
        if not any(cell.strip() for cell in raw):
            if title:
                sections[title] = (headers, rows)
            title, headers, rows = None, None, []
            continue
        if title is None:
            title = raw[0]
        elif headers is None:
            headers = raw
        else:
            rows.append(raw)
    if title:
        sections[title] = (headers, rows)
    return sections


def _parse_excel_sections(response):
    """Pull sheets back out of a rendered workbook: {sheet: (headers, rows)}."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(response.content))
    out = {}
    for ws in wb.worksheets:
        values = [
            ["" if c is None else str(c) for c in row]
            for row in ws.iter_rows(values_only=True)
        ]
        out[ws.title] = (values[0], values[1:]) if values else ([], [])
    return out


class ExportParityTests(TestCase):
    """The CSV and the Excel download must carry the same data."""

    def setUp(self):
        self.user = User.objects.create_user("p@x.com", "Admin", role="admin")
        room = Room.objects.create(room_name="Lab 1", room_type="lab")
        low = Item.objects.create(
            item_name="Cable", unit="pcs", quantity=1, min_quantity=10,
            opening_quantity=1, room=room,
        )
        Item.objects.create(
            item_name="Chair", unit="pcs", quantity=40, min_quantity=5,
            opening_quantity=40, room=room,
        )
        StockTransaction.objects.create(item=low, type="IN", quantity=1, user=self.user)

    def test_both_exports_expose_the_same_sections(self):
        sections = build_report(90)
        csv_blocks = _parse_csv_sections(exports.build_csv(90))
        xlsx_sheets = _parse_excel_sections(exports.build_excel(90))

        self.assertEqual(len(csv_blocks), len(sections))
        self.assertEqual(len(xlsx_sheets), len(sections))
        self.assertEqual({s.title for s in sections}, set(csv_blocks))
        self.assertEqual({s.key for s in sections}, set(xlsx_sheets))

    def test_both_renderers_emit_identical_headers_and_rows(self):
        # Render both from ONE report: build_report() stamps the current time
        # into Summary, so calling it per-format would race the clock.
        sections = build_report(90)
        csv_blocks = _parse_csv_sections(exports.render_csv(sections, filename="a.csv"))
        xlsx_sheets = _parse_excel_sections(exports.render_excel(sections, filename="a.xlsx"))

        for section in sections:
            csv_headers, csv_rows = csv_blocks[section.title]
            xl_headers, xl_rows = xlsx_sheets[section.key]

            self.assertEqual(csv_headers, list(section.headers), section.key)
            self.assertEqual(xl_headers, list(section.headers), section.key)

            expected = [[str(v) for v in row] for row in section.rows] or [["(no records)"]]
            self.assertEqual(csv_rows, expected, section.key)
            self.assertEqual([r[: len(expected[0])] for r in xl_rows], expected, section.key)

    def test_public_builders_agree_on_row_counts(self):
        """build_csv / build_excel go through separate build_report() calls, so
        compare shape rather than the volatile generated-at stamp."""
        csv_blocks = _parse_csv_sections(exports.build_csv(90))
        xlsx_sheets = _parse_excel_sections(exports.build_excel(90))

        for section in build_report(90):
            _csv_headers, csv_rows = csv_blocks[section.title]
            _xl_headers, xl_rows = xlsx_sheets[section.key]
            self.assertEqual(len(csv_rows), len(xl_rows), section.key)

    def test_report_includes_the_sections_that_used_to_be_one_sided(self):
        keys = {s.key for s in build_report(90)}
        self.assertIn("By Category", keys)        # was CSV-only
        self.assertIn("Requisition Items", keys)  # was CSV-only
        self.assertIn("Summary", keys)            # was Excel-only

    def test_snapshot_exports_also_match(self):
        csv_blocks = _parse_csv_sections(exports.build_snapshot_csv())
        xlsx_sheets = _parse_excel_sections(exports.build_snapshot_excel())
        self.assertEqual(len(csv_blocks), len(xlsx_sheets))
        self.assertIn("ITEMS BY ROOM", csv_blocks)
        self.assertIn("By Room", xlsx_sheets)


class LowStockListTests(TestCase):
    def setUp(self):
        self.room = Room.objects.create(room_name="Lab 2", room_type="lab")
        Item.objects.create(item_name="Low", unit="pcs", quantity=2, min_quantity=10, room=self.room)
        Item.objects.create(item_name="Empty", unit="pcs", quantity=0, min_quantity=5)
        Item.objects.create(item_name="Fine", unit="pcs", quantity=99, min_quantity=5, room=self.room)

    def test_lists_only_low_items_with_their_room(self):
        rows = low_stock_items()
        self.assertEqual([r["item_name"] for r in rows], ["Low", "Empty"])
        self.assertEqual(rows[0]["room_name"], "Lab 2")
        self.assertEqual(rows[1]["room_name"], "Unassigned")

    def test_worst_shortage_comes_first(self):
        rows = low_stock_items()
        self.assertEqual(rows[0]["shortage"], 8)
        self.assertEqual(rows[1]["shortage"], 5)

    def test_out_of_stock_is_flagged(self):
        rows = {r["item_name"]: r for r in low_stock_items()}
        self.assertFalse(rows["Low"]["out_of_stock"])
        self.assertTrue(rows["Empty"]["out_of_stock"])

    def test_list_length_matches_the_dashboard_count(self):
        self.assertEqual(len(low_stock_items()), dashboard_stats()["low_stock_items"])


class RoomSnapshotTests(TestCase):
    def setUp(self):
        self.lab = Room.objects.create(room_name="Lab 3", room_type="lab")
        Item.objects.create(item_name="A", unit="pcs", quantity=10, min_quantity=1, room=self.lab)
        Item.objects.create(item_name="B", unit="pcs", quantity=2, min_quantity=5, room=self.lab)
        Item.objects.create(item_name="C", unit="pcs", quantity=7, min_quantity=1)

    def test_groups_items_by_room_with_totals(self):
        rooms = {r["room_name"]: r for r in room_snapshot()["rooms"]}
        self.assertEqual(rooms["Lab 3"]["item_count"], 2)
        self.assertEqual(rooms["Lab 3"]["total_quantity"], 12)
        self.assertEqual(rooms["Lab 3"]["low_stock_count"], 1)

    def test_unassigned_items_get_their_own_bucket(self):
        rooms = {r["room_name"]: r for r in room_snapshot()["rooms"]}
        self.assertIn("General Storage", rooms)
        self.assertEqual(rooms["General Storage"]["total_quantity"], 7)

    def test_grand_totals_add_up(self):
        snap = room_snapshot()
        self.assertEqual(snap["item_count"], 3)
        self.assertEqual(snap["total_quantity"], 19)
        self.assertEqual(snap["low_stock_count"], 1)
        self.assertIsNotNone(snap["generated_at"])


class StockMovementSeriesTests(TestCase):
    def setUp(self):
        self.item = Item.objects.create(item_name="W", unit="pcs", quantity=50, opening_quantity=50)

    def _txn(self, txn_type, qty, days_ago=0):
        txn = StockTransaction.objects.create(item=self.item, type=txn_type, quantity=qty)
        if days_ago:
            StockTransaction.objects.filter(pk=txn.pk).update(
                timestamp=timezone.now() - timedelta(days=days_ago)
            )
        return txn

    def test_every_day_in_the_window_is_present(self):
        series = stock_movement_series(30)
        self.assertEqual(len(series["labels"]), 30)
        self.assertEqual(len(series["stock_in"]), 30)
        self.assertEqual(len(series["stock_out"]), 30)

    def test_totals_split_in_from_out(self):
        self._txn("IN", 10)
        self._txn("OUT", 4)
        series = stock_movement_series(30)
        self.assertEqual(series["total_in"], 10)
        self.assertEqual(series["total_out"], 4)
        self.assertEqual(series["net"], 6)
        self.assertEqual(series["stock_in"][-1], 10)
        self.assertEqual(series["stock_out"][-1], 4)

    def test_transactions_outside_the_window_are_excluded(self):
        self._txn("IN", 99, days_ago=45)
        self.assertEqual(stock_movement_series(30)["total_in"], 0)
        self.assertEqual(stock_movement_series(60)["total_in"], 99)

    def test_adjustments_do_not_count_as_movement(self):
        self._txn("ADJUST", 25)
        series = stock_movement_series(30)
        self.assertEqual(series["total_in"], 0)
        self.assertEqual(series["total_out"], 0)

    def test_day_count_is_clamped_to_a_sane_range(self):
        self.assertEqual(stock_movement_series(0)["days"], 1)
        self.assertEqual(stock_movement_series(9999)["days"], 365)
